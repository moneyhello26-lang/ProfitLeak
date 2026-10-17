import csv
import hashlib
import io
import json
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from pathlib import PureWindowsPath
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from defusedxml.common import DefusedXmlException
from openpyxl import load_workbook
from pydantic import ValidationError

from app.schemas.imports import (
    ExpensesRow,
    ImportKind,
    PreviewRow,
    ProductsRow,
    RowError,
    SalesRow,
)

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_ROWS = 20_000
MAX_COLUMNS = 32
MAX_ERRORS = 100
MAX_EXPANDED_BYTES = 50 * 1024 * 1024
ROW_MODELS = {"sales": SalesRow, "expenses": ExpensesRow, "products": ProductsRow}


class FileProblem(ValueError):
    def __init__(self, code: str, message: str, status_code: int = 400):
        super().__init__(message)
        self.code = code
        self.status_code = status_code


@dataclass
class ParsedFile:
    filename: str
    file_sha256: str
    rows: list[PreviewRow] = field(default_factory=list)
    errors: list[RowError] = field(default_factory=list)
    row_count: int = 0
    error_count: int = 0

    def error(self, row: int, name: str, message: str):
        self.error_count += 1
        if len(self.errors) < MAX_ERRORS:
            self.errors.append(RowError(row=row, field=name, message=message))

    @property
    def content_sha256(self) -> str:
        # Ignore row order, formatting, and numeric trailing zeroes; retain repeated rows.
        def canonical(value):
            if isinstance(value, Decimal):
                return "0" if value == 0 else format(value.normalize(), "f")
            if isinstance(value, (date, datetime)):
                return value.isoformat()
            return value

        rows = [
            json.dumps(
                {key: canonical(value) for key, value in row.data.model_dump().items()},
                sort_keys=True,
                ensure_ascii=False,
            )
            for row in self.rows
        ]
        return hashlib.sha256("\n".join(sorted(rows)).encode("utf-8")).hexdigest()


def parse_file(filename: str, data: bytes, kind: ImportKind, *, model=None, row_type=PreviewRow):
    # Windows and POSIX path separators are both stripped from client filenames.
    filename = PureWindowsPath(filename).name
    if not filename or len(filename) > 255:
        raise FileProblem("invalid_filename", "Некорректное имя файла.")
    if len(data) > MAX_UPLOAD_BYTES:
        raise FileProblem("file_too_large", "Файл превышает лимит 10 МБ.", 413)
    result = ParsedFile(filename=filename, file_sha256=hashlib.sha256(data).hexdigest())
    extension = PureWindowsPath(filename).suffix.lower()
    if extension == ".csv":
        parse_csv(data, kind, result, model=model, row_type=row_type)
    elif extension == ".xlsx":
        parse_xlsx(data, kind, result, model=model, row_type=row_type)
    else:
        raise FileProblem("unsupported_format", "Используйте CSV или XLSX.", 415)
    if result.row_count == 0 and not result.errors:
        result.error(1, "file", "Файл не содержит строк данных.")
    return result


def consume_rows(rows, kind: ImportKind, result: ParsedFile, *, model=None, row_type=PreviewRow):
    model = model or ROW_MODELS[kind]
    header = next(rows, None)
    if header is None:
        result.error(1, "file", "Файл пуст.")
        return
    names = [str(value).strip() if value is not None else "" for value in header]
    if len(names) > MAX_COLUMNS or len(names) != len(set(names)) or "" in names:
        result.error(1, "columns", "Заголовки должны быть уникальными и непустыми.")
        return
    missing = [
        name
        for name, info in model.model_fields.items()
        if info.is_required() and name not in names
    ]
    extra = [name for name in names if name not in model.model_fields]
    for name in missing:
        result.error(1, name, "Отсутствует обязательная колонка.")
    for name in extra:
        result.error(1, name, "Неизвестная колонка.")
    if missing or extra:
        return
    product_names = {}
    for number, values in enumerate(rows, start=2):
        if number > MAX_ROWS + 1:
            raise FileProblem("too_many_rows", "Файл превышает лимит 20 000 строк.", 413)
        if all(value is None or value == "" for value in values):
            continue
        result.row_count += 1
        if len(values) != len(names):
            result.error(number, "columns", "Число значений не совпадает с числом колонок.")
            continue
        # An XLSX formula is represented separately; cached values are never imported.
        if any(isinstance(value, Formula) for value in values):
            for name, value in zip(names, values, strict=True):
                if isinstance(value, Formula):
                    result.error(number, name, "Формулы XLSX запрещены. Вставьте значения.")
            continue
        record = dict(zip(names, values, strict=True))
        try:
            validated = model.model_validate(record)
        except ValidationError as error:
            for detail in error.errors():
                result.error(number, str(detail["loc"][0]), validation_message(detail["type"]))
            continue
        if kind in ("products", "sales"):
            name = validated.name if kind == "products" else validated.product_name
            previous = product_names.get(validated.product_id)
            if previous is not None and (kind == "products" or previous != name):
                result.error(
                    number, "product_id", "Повторный ID товара или несовпадающие названия."
                )
                continue
            product_names[validated.product_id] = name
        result.rows.append(row_type(source_row=number, data=validated))


def validation_message(error_type: str) -> str:
    if error_type in ("decimal_max_digits", "decimal_max_places", "decimal_whole_digits"):
        return "Слишком много цифр или знаков после запятой."
    if error_type in ("greater_than", "greater_than_equal"):
        return "Количество должно быть положительным; суммы не могут быть отрицательными."
    if error_type == "literal_error":
        return "Недопустимый тип записи."
    return "Некорректное значение. Проверьте формат, длину и обязательность поля."


def parse_csv(data: bytes, kind: ImportKind, result: ParsedFile, **options):
    try:
        text = data.decode("utf-8-sig")
        if "\x00" in text:
            raise FileProblem("invalid_csv", "Файл CSV содержит недопустимые символы.")
        consume_rows(iter(csv.reader(io.StringIO(text, newline=""), strict=True)),
                     kind, result, **options)
    except UnicodeDecodeError as error:
        raise FileProblem("invalid_encoding", "CSV должен быть в кодировке UTF-8.") from error
    except csv.Error as error:
        raise FileProblem("invalid_csv", "Не удалось прочитать CSV.") from error


class Formula:
    pass


def parse_xlsx(data: bytes, kind: ImportKind, result: ParsedFile, **options):
    workbook = None
    try:
        with ZipFile(io.BytesIO(data)) as archive:
            if sum(entry.file_size for entry in archive.infolist()) > MAX_EXPANDED_BYTES:
                raise FileProblem("expanded_file_too_large", "Содержимое XLSX слишком велико.", 413)
            if len(archive.infolist()) > 1000:
                raise FileProblem("invalid_xlsx", "Слишком много элементов в XLSX.")
            if "xl/workbook.xml" not in archive.namelist():
                raise FileProblem("invalid_xlsx", "Файл не является книгой XLSX.")
        workbook = load_workbook(
            io.BytesIO(data), read_only=True, data_only=False, keep_links=False
        )
        if len(workbook.worksheets) != 1:
            raise FileProblem("multiple_sheets", "Оставьте в XLSX один лист с данными.")
        sheet = workbook.worksheets[0]
        if sheet.max_column is not None and sheet.max_column > MAX_COLUMNS:
            raise FileProblem("too_many_columns", "В файле слишком много колонок.")
        # Do not trust worksheet dimensions; iterate actual records under row/column limits.
        sheet.reset_dimensions()

        def rows():
            width = None
            for cells in sheet.iter_rows(max_col=MAX_COLUMNS + 1):
                values = [Formula() if cell.data_type == "f" else cell.value for cell in cells]
                while values and values[-1] is None:
                    values.pop()
                if width is None:
                    width = len(values)
                elif len(values) < width:
                    values.extend([None] * (width - len(values)))
                yield values

        consume_rows(iter(rows()), kind, result, **options)
    except FileProblem:
        raise
    except (
        BadZipFile,
        ValueError,
        KeyError,
        OSError,
        EOFError,
        ParseError,
        DefusedXmlException,
    ) as error:
        raise FileProblem("invalid_xlsx", "Не удалось прочитать XLSX.") from error
    finally:
        if workbook is not None:
            workbook.close()
